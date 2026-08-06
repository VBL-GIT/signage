# -*- coding: utf-8 -*-
"""Generate ready-to-upload demo Excel files (stores, users, tasks) for the live demo.
Upload order: stores -> users -> tasks (tasks reference the stores below)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = os.path.dirname(os.path.abspath(__file__))
HEADER_FILL = PatternFill("solid", start_color="0B5394")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="C9D4E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(filename, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for r in rows:
        ws.append(r)
    for ri in range(2, 2 + len(rows)):
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
    for ci, h in enumerate(headers, start=1):
        maxlen = max([len(str(h))] + [len(str(row[ci - 1])) if ci - 1 < len(row) else 0 for row in rows])
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(max(maxlen + 3, 12), 40)
    ws.freeze_panes = "A2"
    path = os.path.join(OUT, filename)
    wb.save(path)
    print("WROTE", path)


# ---------------------------------------------------------------- 1) STORES
build(
    "demo_stores.xlsx",
    ["name", "address", "pincode", "lat", "long", "uid", "contact_no"],
    [
        ["VBL Outlet Andheri", "Plot 12, Andheri West, Mumbai", "400053", 19.1197, 72.8468, "VBL-STR-101", "9820010101"],
        ["VBL Outlet CP", "N-Block, Connaught Place, New Delhi", "110001", 28.6315, 77.2167, "VBL-STR-102", "9820010102"],
        ["VBL Outlet Koramangala", "80 Ft Road, Koramangala, Bengaluru", "560034", 12.9352, 77.6245, "VBL-STR-103", "9820010103"],
        ["VBL Outlet Banjara Hills", "Road No 12, Banjara Hills, Hyderabad", "500034", 17.4156, 78.4347, "VBL-STR-104", "9820010104"],
        ["VBL Outlet Salt Lake", "Sector V, Salt Lake, Kolkata", "700091", 22.5808, 88.4172, "VBL-STR-105", "9820010105"],
        ["VBL Outlet Aundh", "ITI Road, Aundh, Pune", "411007", 18.5590, 73.8076, "VBL-STR-106", "9820010106"],
    ],
)

# ---------------------------------------------------------------- 2) USERS
build(
    "demo_users.xlsx",
    ["first_name", "last_name", "email", "role", "mobile", "password", "vendor_uid"],
    [
        ["Rohit", "Mehra", "rohit.mehra@acme.demo", "employee", "9820020001", "", "VND-001"],
        ["Priya", "Nair", "priya.nair@acme.demo", "employee", "9820020002", "", "VND-001"],
        ["Karan", "Singh", "karan.singh@acme.demo", "employee", "9820020003", "", "VND-001"],
        ["Anjali", "Rao", "anjali.rao@acme.demo", "vendor_user", "9820020004", "", "VND-001"],
        ["Vikram", "Desai", "vikram.desai@rjcorp.demo", "rjcorp_user", "9820020005", "", ""],
    ],
)

# ---------------------------------------------------------------- 3) TASKS
# store_name values MUST match the stores uploaded above (case-insensitive).
build(
    "demo_tasks.xlsx",
    ["task_type", "installation_type", "vendor_uid", "store_name", "pincode",
     "target_pamphlet_count", "brand_name", "boarding_size_label"],
    [
        ["recee", "", "VND-001", "VBL Outlet Andheri", "", "", "", ""],
        ["recee", "", "VND-001", "VBL Outlet Koramangala", "", "", "", ""],
        ["installation", "direct", "VND-001", "VBL Outlet CP", "110001", 300, "", ""],
        ["installation", "direct", "VND-001", "VBL Outlet Salt Lake", "700091", 150, "", ""],
        ["installation", "direct_boarding", "VND-001", "VBL Outlet Banjara Hills", "", "", "BrandX", "Medium (4x3 ft)"],
        ["installation", "direct_boarding", "VND-001", "VBL Outlet Aundh", "", "", "BrandY", "Large (6x4 ft)"],
    ],
)
